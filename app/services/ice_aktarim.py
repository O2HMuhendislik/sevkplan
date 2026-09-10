"""Excel içe aktarım servisleri: ürün master datası ve siparişler."""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.domain.iller import yer_adi
from app.domain.metin import buyuk_harf
from app.models import (
    IceAktarim,
    Musteri,
    SiparisDurumu,
    SiparisSatiri,
    NakliyeTarifesi,
    Urun,
    UrunBagi,
)
from app.db import parcali_scalars
from app.services import excel
from app.services.excel import ExcelHatasi
from app.services.veri_formatlari import (
    MUSTERI_ALANLARI,
    MUSTERI_ALIAS,
    BUTCE_ALANLARI,
    BUTCE_ALIAS,
    EK_UCRET_ALANLARI,
    EK_UCRET_ALIAS,
    TARIFE_ALANLARI,
    TARIFE_ALIAS,
    URUN_BAGI_ALANLARI,
    URUN_BAGI_ALIAS,
    SIPARIS_ALANLARI,
    SIPARIS_ALIAS,
    URUN_ALANLARI,
    URUN_ALIAS,
    bayi_adini_coz,
    yer_alanlarini_coz,
    zorunlu_alanlar,
)


RING_ILI = "ESKISEHIR"
"""Ring, Eskişehir içi dağıtımdır; başka ile giden sipariş bu havuza alınmaz."""


def ring_ili_mi(sehir: object) -> bool:
    """Şehir alanı Ring kapsamında mı?

    Türkçe karakter ve yazım farkları sadeleştirilerek karşılaştırılır. **Boş şehir
    reddedilmez:** bir kısım kaynak dosyada bu sütun hiç doldurulmuyor, depo kodu
    zaten Eskişehir'i gösteriyor. Yalnızca açıkça başka bir il yazılmışsa satır
    alınmaz.
    """
    ad = yer_adi(sehir)
    return not ad or ad == RING_ILI


class RingKapsamDisi(ExcelHatasi):
    """Ring havuzuna Eskişehir dışı bir il ile sipariş yüklenmeye çalışıldı."""

    def __init__(self, sehir: str) -> None:
        super().__init__(
            f"Ring yalnızca Eskişehir içi dağıtımdır; '{sehir}' ili alınmadı. "
            "Bu siparişi İç Piyasa modülünden yükleyin."
        )


@dataclass
class SatirHatasi:
    satir_no: int
    anahtar: str
    mesaj: str


@dataclass
class IceAktarimSonucu:
    toplam: int = 0
    eklenen: int = 0
    guncellenen: int = 0
    atlanan: int = 0
    birlestirilen: int = 0
    """Aynı sipariş/teslimat/ürün tekrar geldiği için miktarı toplanan satır sayısı."""
    hatalar: list[SatirHatasi] = field(default_factory=list)
    uyarilar: list[SatirHatasi] = field(default_factory=list)
    """Kayıt alındı ama eksik veri var; kullanıcının görmesi gereken durumlar."""
    reddedilen: int = 0
    """Modülün kapsamı dışında olduğu için hiç alınmayan satır sayısı."""

    @property
    def basarili(self) -> int:
        return self.eklenen + self.guncellenen

    @property
    def hatali(self) -> int:
        return len(self.hatalar)

    def ozet(self) -> str:
        metin = (
            f"{self.toplam} satır okundu · {self.eklenen} yeni · "
            f"{self.guncellenen} güncellendi · {self.atlanan} atlandı · "
            f"{self.hatali} hatalı"
        )
        if self.reddedilen:
            metin += f" · {self.reddedilen} satır modül kapsamı dışında, alınmadı"
        if self.birlestirilen:
            metin += f" · {self.birlestirilen} satır birleştirildi"
        if self.uyarilar:
            metin += f" · {len(self.uyarilar)} eksik veri uyarısı"
        return metin


def _kontrol_et(dosya: Any, alanlar, alias) -> None:
    eksikler = excel.eksik_kolonlar(dosya, alias, zorunlu_alanlar(alanlar))
    if eksikler:
        basliklar = {alan.ad: alan.baslik for alan in alanlar}
        raise ExcelHatasi(
            "Dosyada zorunlu kolonlar bulunamadı: "
            + ", ".join(basliklar[ad] for ad in eksikler)
        )


def _doluyu_yaz(kayit, alan: str, deger) -> None:
    """Yalnızca **dolu** değeri yazar; boş gelen sütun mevcut veriyi silmez.

    Master data dosyaları çoğu zaman kısmidir: kullanıcı listeyi indirip yalnızca
    eksik ölçüleri dolduruyor, ya da elindeki dosyada sütunların bir kısmı yok. Boş
    hücreyi "sil" diye okursak tek bir kısmi yükleme bütün ölçüleri siler ve
    planlama durur.

    Bir alanı **kasten** boşaltmak Master Data ekranındaki tekil düzenleme formuyla
    yapılır; orada boş bırakmak silme anlamına gelir.
    """
    if deger is None or deger == "":
        return
    setattr(kayit, alan, deger)


def urunleri_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    _kontrol_et(dosya, URUN_ALANLARI, URUN_ALIAS)
    kayitlar = excel.satirlari_oku(dosya, URUN_ALIAS, zorunlu_alanlar(URUN_ALANLARI))
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))

    mevcutlar = {
        urun.urun_kodu: urun for urun in db.scalars(select(Urun)).all()
    }
    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        urun_kodu = excel.metin(kayit.get("urun_kodu"))
        try:
            if not urun_kodu:
                raise ExcelHatasi("Stok kodu boş olamaz")
            urun_adi = excel.metin(kayit.get("urun_adi"))
            if not urun_adi:
                raise ExcelHatasi("Stok adı boş olamaz")
            palet_ici = excel.tam_sayi_ya_da(kayit.get("palet_ici_adet"))
            kamyon_adet = excel.tam_sayi_ya_da(kayit.get("kamyon_yukleme_adeti"))
            tir_adet = excel.tam_sayi_ya_da(kayit.get("tir_yukleme_adeti"))
        except ExcelHatasi as hata:
            sonuc.hatalar.append(SatirHatasi(satir_no, urun_kodu or "-", str(hata)))
            continue

        urun = mevcutlar.get(urun_kodu)
        if urun is None:
            urun = Urun(urun_kodu=urun_kodu)
            db.add(urun)
            mevcutlar[urun_kodu] = urun
            sonuc.eklenen += 1
        else:
            sonuc.guncellenen += 1

        urun.urun_adi = urun_adi
        # Türkçe büyük harf: 'Klima' -> 'KLİMA'. Python'un upper()'ı 'KLIMA' verip
        # aynı grubu ikiye bölüyordu (bkz. app/domain/metin.py).
        _doluyu_yaz(urun, "urun_grubu", buyuk_harf(kayit.get("urun_grubu")) or None)
        _doluyu_yaz(urun, "palet_ici_adet", palet_ici)
        _doluyu_yaz(urun, "kamyon_yukleme_adeti", kamyon_adet)
        _doluyu_yaz(urun, "kamyon_palet", excel.tam_sayi_ya_da(kayit.get("kamyon_palet")))
        _doluyu_yaz(urun, "tir_yukleme_adeti", tir_adet)
        _doluyu_yaz(urun, "tir_palet", excel.tam_sayi_ya_da(kayit.get("tir_palet")))
        _doluyu_yaz(urun, "agirlik", excel.sayi_ya_da(kayit.get("agirlik")))
        _doluyu_yaz(urun, "desi", excel.sayi_ya_da(kayit.get("desi")))
        _doluyu_yaz(urun, "m3", excel.sayi_ya_da(kayit.get("m3")))
        _doluyu_yaz(urun, "palet_en", excel.tam_sayi_ya_da(kayit.get("palet_en")))
        _doluyu_yaz(urun, "palet_boy", excel.tam_sayi_ya_da(kayit.get("palet_boy")))
        _doluyu_yaz(urun, "palet_yukseklik",
                    excel.tam_sayi_ya_da(kayit.get("palet_yukseklik")))
        _doluyu_yaz(urun, "header_kod", excel.metin(kayit.get("header_kod")))
        if not excel.bos_mu(kayit.get("aktif")):
            urun.aktif = excel.evet_hayir(kayit.get("aktif"), True)

        if not urun.planlanabilir_mi:
            # Kayıt yine de alınır; planlamaya girdiğinde gerekçesiyle uyarılır.
            sonuc.uyarilar.append(
                SatirHatasi(
                    satir_no, urun_kodu,
                    "Palet içi adet / kamyon / tır yükleme adeti alanlarının üçü de boş; "
                    "bu ürün planlamaya giremez",
                )
            )

    _aktarim_kaydet(db, dosya_adi, "URUN", sonuc, kullanici)
    db.flush()
    return sonuc


def _tanitici_alanlari_tazele(satir: SiparisSatiri, kayit: dict) -> None:
    """Planı etkilemeyen tanıtıcı alanları dosyadaki güncel değerle doldurur.

    Yalnızca **dolu** bir değer yazılır; dosyada boş gelen sütun sistemdeki mevcut
    bilgiyi silmez. Planlanmış satırlarda da çalıştığı için bu güvence şart.
    """
    bayi = excel.metin(kayit.get("bayi_adi")) or excel.metin(kayit.get("ikinci_not"))
    if bayi:
        satir.bayi_adi = bayi
    firma, adres, ilce, incoterms = yer_alanlarini_coz(
        excel.metin(kayit.get("alici_firma")),
        excel.metin(kayit.get("sevk_adresi")),
        excel.metin(kayit.get("teslim_sekli")),
    )
    if firma:
        satir.alici_firma = firma
    if adres:
        satir.sevk_adresi = adres
    if ilce:
        satir.ilce = ilce
    if incoterms:
        satir.incoterms = incoterms


def _mevcut_siparis_satirlari(db: Session, kayitlar: list[dict]) -> dict:
    """Dosyadaki anahtarlara karşılık gelen mevcut satırları **tek seferde** okur.

    Önceden her satır için ayrı bir SELECT atılıyordu: 139 bin satırlık bir dosyada
    139 bin sorgu demek. Hem dakikalar sürüyordu hem de her sorgunun döndürdüğü
    nesne oturumun kimlik haritasında kalıcı olarak birikip belleği şişiriyordu —
    büyük dosyalarda yükleme bu yüzden çöküyordu.

    Yalnızca dosyada geçen sipariş numaraları okunur; bütün tablo belleğe alınmaz.
    """
    siparis_numaralari = {
        excel.metin(kayit.get("siparis_no"))
        for kayit in kayitlar
        if excel.metin(kayit.get("siparis_no"))
    }
    mevcutlar: dict[tuple[str, str, str], SiparisSatiri] = {}
    for satir in parcali_scalars(
        db,
        lambda parca: select(SiparisSatiri).where(SiparisSatiri.siparis_no.in_(parca)),
        sorted(siparis_numaralari),
    ):
        mevcutlar[(satir.siparis_no, satir.teslimat_no, satir.siparis_satir_no)] = satir
    return mevcutlar


def siparisleri_aktar(
    db: Session,
    dosya: Path | Any,
    dosya_adi: str,
    kullanici: str = "sistem",
    modul: str = "RING",
) -> IceAktarimSonucu:
    """Sipariş dosyasını yükler.

    `modul` siparişi hangi havuza yazacağımızı söyler: RING, ROTA (iç piyasa) ya da
    IHRACAT. Her modül yalnızca kendi havuzunu görür ve planlar; aynı satır iki
    modülde birden görünmez.
    """
    _kontrol_et(dosya, SIPARIS_ALANLARI, SIPARIS_ALIAS)
    kayitlar = excel.satirlari_oku(dosya, SIPARIS_ALIAS, zorunlu_alanlar(SIPARIS_ALANLARI))
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))
    aktarim = _aktarim_kaydet(db, dosya_adi, f"SIPARIS/{modul}", sonuc, kullanici)
    db.flush()

    etkilenen_teslimatlar: set[str] = set()
    parti: dict[tuple[str, str, str], SiparisSatiri] = {}
    """Aynı dosyada tekrar eden (sipariş, teslimat, ürün) satırları birleştirmek için."""

    mevcutlar = _mevcut_siparis_satirlari(db, kayitlar)

    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        siparis_no = excel.metin(kayit.get("siparis_no"))
        urun_kodu = excel.metin(kayit.get("urun_kodu"))
        # Kaynak dosyalarda satır numarası yok; satır anahtarı ürün kodudur.
        siparis_satir_no = excel.metin(kayit.get("siparis_satir_no")) or urun_kodu
        anahtar = f"{siparis_no or '-'}/{siparis_satir_no or '-'}"
        try:
            if not siparis_no:
                raise ExcelHatasi("Sipariş No boş olamaz")
            if not urun_kodu:
                raise ExcelHatasi("Stok kodu boş olamaz")
            teslimat_no = excel.metin(kayit.get("teslimat_no"))
            if not teslimat_no:
                raise ExcelHatasi("Teslimat No boş olamaz")
            if not any(karakter.isdigit() for karakter in str(teslimat_no)):
                # Bayi ortak deposu (-1) satırlarında bu sütun teslimat numarası yerine
                # "BAYİ DEPO" gibi bir etiket taşıyor. Sipariş bölünemez birim olduğu
                # için teslimat anahtarı olarak sipariş numarası kullanılır.
                teslimat_no = f"{siparis_no}-{yer_adi(teslimat_no) or 'SIPARIS'}"
            depo_kodu = excel.metin(kayit.get("depo_kodu"))
            # "-1" bayi ortak deposudur (Eskişehir) — geçerli bir depo kodudur.
            if not depo_kodu or depo_kodu.strip() == "0":
                raise ExcelHatasi(
                    f"Depo kodu atanmamış ({depo_kodu or 'boş'}); bu satır planlanamaz"
                )
            sehir = excel.metin(kayit.get("sehir"))
            if modul == "RING" and not ring_ili_mi(sehir):
                # Ring, Eskişehir içi dağıtımdır. Şehir dışı sipariş buraya yüklenirse
                # planlamaya girer ve sahada yanlış araca binerdi; alınmaz.
                raise RingKapsamDisi(sehir or "boş")
            miktar = excel.sayi(kayit.get("miktar"), "Miktar")
            if miktar <= 0:
                raise ExcelHatasi("Miktar sıfırdan büyük olmalı")
            siparis_tarihi = excel.tarih(kayit.get("siparis_tarihi"))
            termin_tarihi = excel.tarih(kayit.get("termin_tarihi"))
        except RingKapsamDisi as hata:
            sonuc.reddedilen += 1
            sonuc.hatalar.append(SatirHatasi(satir_no, anahtar, str(hata)))
            continue
        except ExcelHatasi as hata:
            sonuc.hatalar.append(SatirHatasi(satir_no, anahtar, str(hata)))
            continue

        satir_anahtari = (siparis_no, teslimat_no, siparis_satir_no)
        if satir_anahtari in parti:
            # Kaynak dosyada aynı sipariş/teslimat/ürün birden çok satırda gelebiliyor;
            # miktarlar toplanır, mükerrer kayıt oluşmaz.
            parti[satir_anahtari].miktar = Decimal(parti[satir_anahtari].miktar) + miktar
            sonuc.birlestirilen += 1
            continue

        mevcut = mevcutlar.get(satir_anahtari)
        if mevcut is not None and mevcut.durum in {
            SiparisDurumu.PLANLANDI,
            SiparisDurumu.TAMAMLANDI,
        }:
            # Planlanmış satır yeniden yüklemeyle bozulmaz: miktar, depo, durum ve
            # plan bağı olduğu gibi kalır. Yalnızca **tanıtıcı** alanlar tazelenir —
            # bayi adı, alıcı, adres, ilçe. Bunlar planı değiştirmiyor ama yükleme
            # formunda görünüyor; boş kalırsa depo formu okuyamıyor ve satır zaten
            # planlandığı için bir daha hiç düzelmiyordu.
            _tanitici_alanlari_tazele(mevcut, kayit)
            sonuc.atlanan += 1
            continue

        if mevcut is None:
            mevcut = SiparisSatiri(
                siparis_no=siparis_no,
                teslimat_no=teslimat_no,
                siparis_satir_no=siparis_satir_no,
            )
            db.add(mevcut)
            mevcutlar[satir_anahtari] = mevcut
            sonuc.eklenen += 1
        else:
            sonuc.guncellenen += 1

        mevcut.teslimat_no = teslimat_no
        mevcut.urun_kodu = urun_kodu
        mevcut.urun_adi = excel.metin(kayit.get("urun_adi"))
        mevcut.miktar = miktar
        mevcut.depo_kodu = depo_kodu.strip().upper()
        mevcut.sehir = sehir
        # BayiAdi bazı dosyalarda boş geliyor; gerçek ad ikinci `Not` sütununda.
        mevcut.bayi_adi = excel.metin(kayit.get("bayi_adi")) or excel.metin(
            kayit.get("ikinci_not")
        )
        mevcut.teslim_sekli = excel.metin(kayit.get("teslim_sekli"))
        # Alıcı firma / adres / ilçe sütunlarının anlamı satır tipine göre kayıyor;
        # hangisinin ne olduğu içeriğe bakılarak çözülür (bkz. yer_alanlarini_coz).
        firma, adres, ilce, incoterms = yer_alanlarini_coz(
            excel.metin(kayit.get("alici_firma")),
            excel.metin(kayit.get("sevk_adresi")),
            mevcut.teslim_sekli,
        )
        mevcut.alici_firma = firma or None
        mevcut.sevk_adresi = adres or None
        mevcut.incoterms = incoterms or None
        mevcut.ilce = ilce or None
        mevcut.siparis_tarihi = siparis_tarihi
        mevcut.termin_tarihi = termin_tarihi
        mevcut.durum = SiparisDurumu.BEKLEMEDE
        mevcut.hata_aciklamasi = None
        mevcut.bekleme_sebebi = None
        mevcut.modul = modul
        mevcut.ice_aktarim_id = aktarim.id
        parti[satir_anahtari] = mevcut
        etkilenen_teslimatlar.add(teslimat_no)

    db.flush()
    for hata in _teslimatlari_dogrula(db, etkilenen_teslimatlar):
        sonuc.hatalar.append(hata)

    aktarim.basarili_satir = sonuc.basarili
    aktarim.hatali_satir = sonuc.hatali
    aktarim.hata_ozeti = _hata_ozeti(sonuc)
    db.flush()
    return sonuc


def _teslimatlari_dogrula(db: Session, teslimat_nolar: set[str]) -> list[SatirHatasi]:
    """Bir teslimat tek depoya ait olmak zorunda.

    Teslimatın birden fazla ürün içermesi hata değildir: saf plana giremez ama baskın
    ürün grubunun karma planına yazılır (bkz. app/services/planlama_anahtari.py).
    """
    hatalar: list[SatirHatasi] = []
    if not teslimat_nolar:
        return hatalar

    # Değer listeleri veri büyüdükçe binleri buluyor; parçalanmazsa SQLite
    # "too many SQL variables" ile düşer.
    satirlar = parcali_scalars(
        db,
        lambda parca: select(SiparisSatiri).where(
            SiparisSatiri.teslimat_no.in_(parca)
        ),
        teslimat_nolar,
    )
    urun_haritasi = {
        urun.urun_kodu: urun
        for urun in parcali_scalars(
            db,
            lambda parca: select(Urun).where(Urun.urun_kodu.in_(parca)),
            {satir.urun_kodu for satir in satirlar},
        )
    }

    gruplar: dict[str, list[SiparisSatiri]] = {}
    for satir in satirlar:
        gruplar.setdefault(satir.teslimat_no, []).append(satir)

    for teslimat_no, grup in gruplar.items():
        mesaj = None
        depolar = {satir.depo_kodu for satir in grup}
        if len(depolar) > 1:
            mesaj = (
                f"Teslimat birden fazla depo kodu içeriyor ({', '.join(sorted(depolar))})."
            )
        if mesaj:
            for satir in grup:
                satir.durum = SiparisDurumu.HATALI
                satir.hata_aciklamasi = mesaj
            hatalar.append(SatirHatasi(0, teslimat_no, mesaj))
    return hatalar


def _hata_ozeti(sonuc: IceAktarimSonucu, azami: int = 50) -> str | None:
    if not sonuc.hatalar:
        return None
    satirlar = [
        f"Satır {hata.satir_no or '-'} ({hata.anahtar}): {hata.mesaj}"
        for hata in sonuc.hatalar[:azami]
    ]
    if len(sonuc.hatalar) > azami:
        satirlar.append(f"... ve {len(sonuc.hatalar) - azami} hata daha")
    return "\n".join(satirlar)


def _aktarim_kaydet(
    db: Session, dosya_adi: str, tur: str, sonuc: IceAktarimSonucu, kullanici: str
) -> IceAktarim:
    aktarim = IceAktarim(
        dosya_adi=dosya_adi,
        tur=tur,
        toplam_satir=sonuc.toplam,
        basarili_satir=sonuc.basarili,
        hatali_satir=sonuc.hatali,
        hata_ozeti=_hata_ozeti(sonuc),
        kullanici=kullanici,
    )
    db.add(aktarim)
    return aktarim


def musterileri_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """İç piyasa müşteri master datasını yükler.

    Anahtar bayi adının normalize hâlidir: kaynak dosyalarda aynı bayi hem 'İSTANBUL
    ISITMA' hem 'ISTANBUL ISITMA' geçebiliyor; normalize edilmezse aynı müşteri iki
    ayrı kayda bölünür ve 3 palet kuralı yanlış hesaplanır.
    """
    _kontrol_et(dosya, MUSTERI_ALANLARI, MUSTERI_ALIAS)
    kayitlar = excel.satirlari_oku(
        dosya, MUSTERI_ALIAS, zorunlu_alanlar(MUSTERI_ALANLARI)
    )
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))

    mevcutlar = {m.anahtar: m for m in db.scalars(select(Musteri)).all()}
    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        bayi_adi = excel.metin(kayit.get("bayi_adi"))
        anahtar = yer_adi(bayi_adi)
        if not anahtar:
            sonuc.hatalar.append(SatirHatasi(satir_no, "-", "Bayi adı boş olamaz"))
            continue

        musteri = mevcutlar.get(anahtar)
        if musteri is None:
            musteri = Musteri(anahtar=anahtar, bayi_adi=bayi_adi)
            db.add(musteri)
            mevcutlar[anahtar] = musteri
            sonuc.eklenen += 1
        else:
            sonuc.guncellenen += 1

        # Ürün master datasındaki kural burada da geçerli: boş gelen sütun mevcut
        # bilgiyi silmez (bkz. _doluyu_yaz).
        musteri.bayi_adi = bayi_adi
        _doluyu_yaz(musteri, "bayi_kodu", excel.metin(kayit.get("bayi_kodu")))
        _doluyu_yaz(musteri, "alici_firma", excel.metin(kayit.get("alici_firma")))
        _doluyu_yaz(musteri, "il", yer_adi(kayit.get("il")))
        _doluyu_yaz(musteri, "ilce", yer_adi(kayit.get("ilce")))
        _doluyu_yaz(musteri, "sevk_adresi", excel.metin(kayit.get("sevk_adresi")))
        _doluyu_yaz(musteri, "telefon", excel.metin(kayit.get("telefon")))
        _doluyu_yaz(
            musteri, "incoterms",
            (excel.metin(kayit.get("incoterms")) or "").upper() or None,
        )
        tir = (excel.metin(kayit.get("tir_girisi")) or "").strip().upper()[:1]
        if tir in {"E", "H", "?"}:
            musteri.tir_girisi = tir
        elif musteri.tir_girisi is None:
            musteri.tir_girisi = "?"
        _doluyu_yaz(musteri, "bolge_kodu", excel.metin(kayit.get("bolge_kodu")))
        _doluyu_yaz(musteri, "eposta", excel.metin(kayit.get("eposta")))
        _doluyu_yaz(musteri, "ozel_durum", excel.metin(kayit.get("ozel_durum")))
        # Sevk tipi metni araç tipini, cumartesi ve e-irsaliye bilgisini birlikte
        # taşıyor; dosyada varsa üçü de ondan çözülür (bkz. musteri_ek_bilgi).
        sevk = excel.metin(kayit.get("sevk_tipi"))
        if sevk:
            from app.services.musteri_ek_bilgi import sevk_tipini_coz

            cozum = sevk_tipini_coz(sevk)
            musteri.sevk_tipi = cozum["sevk_tipi"]
            musteri.cumartesi_teslimat = cozum["cumartesi_teslimat"]
            musteri.e_irsaliye = cozum["e_irsaliye"]
            if cozum["tir_girisi"] != "?" and excel.bos_mu(kayit.get("tir_girisi")):
                # Sütun ayrıca doldurulmamışsa sevk tipinden türetilen değer geçerli.
                musteri.tir_girisi = cozum["tir_girisi"]
        else:
            if not excel.bos_mu(kayit.get("cumartesi_teslimat")):
                musteri.cumartesi_teslimat = excel.evet_hayir(
                    kayit.get("cumartesi_teslimat"), True
                )
            if not excel.bos_mu(kayit.get("e_irsaliye")):
                musteri.e_irsaliye = excel.evet_hayir(kayit.get("e_irsaliye"), False)
        _doluyu_yaz(musteri, "notlar", excel.metin(kayit.get("notlar")))
        if not excel.bos_mu(kayit.get("aktif")):
            musteri.aktif = excel.evet_hayir(kayit.get("aktif"), True)

        if not musteri.il:
            sonuc.uyarilar.append(
                SatirHatasi(
                    satir_no, bayi_adi,
                    "İl boş; bu müşteri bölgeye yerleştirilemez ve FTL planına giremez",
                )
            )

    _aktarim_kaydet(db, dosya_adi, "MUSTERI", sonuc, kullanici)
    db.flush()
    return sonuc


def _tonaj(deger: Any) -> Decimal | None:
    """'22.000 KG' / '22 000 KG' / 22000 -> Decimal(22000). Anlaşılmazsa None."""
    ham = excel.metin(deger) or ""
    if not ham:
        return None
    rakamlar = "".join(k for k in ham if k.isdigit())
    if not rakamlar:
        return None
    return Decimal(rakamlar)


def ihracat_urunlerini_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """İhracat ürün master datasını yükler (`Hesaplama.xlsx` → `Ürün` sayfası).

    Şirketin hesaplama dosyası olduğu gibi verilebilir: kolon başlıkları dosyadaki
    adlarla eşleşir. Yükleme adetleri iki sürüm hâlinde saklanır — temel sütunlar
    yeni hesaplama, `-2` sütunları eski hesaplama. Hangisinin kullanılacağını
    müşteri master datası söyler.

    Boş ölçüler hata değildir: master datada 2.862 üründen bir kısmında tır ya da
    konteyner adedi yok. Bunlar uyarı olarak listelenir, planlamada desiden
    yaklaşık hesaplanır.
    """
    from app.models import IhracatUrunu
    from app.services.veri_formatlari import (
        IHRACAT_URUN_ALANLARI,
        IHRACAT_URUN_ALIAS,
    )

    _kontrol_et(dosya, IHRACAT_URUN_ALANLARI, IHRACAT_URUN_ALIAS)
    kayitlar = excel.satirlari_oku(
        dosya, IHRACAT_URUN_ALIAS, zorunlu_alanlar(IHRACAT_URUN_ALANLARI)
    )
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))

    sayisal_alanlar = (
        "palet_ici_adet", "tir_yukleme_adeti", "konteyner_yukleme_adeti",
        "palet_ici_adet_eski", "tir_yukleme_adeti_eski",
        "konteyner_yukleme_adeti_eski", "desi", "agirlik", "dokme_adeti",
    )
    tam_sayi_alanlari = ("en", "boy", "yukseklik")

    # Aynı ürün kodu dosyada birden çok geçebiliyor ve tekrarların bir kısmı boş.
    # Son satır kazanırsa dolu ölçüler silinir; bu yüzden tekrarlar birleştirilir:
    # bir alan bir kez doldurulduysa sonraki boş satır onu ezmez.
    birlesik: dict[str, dict] = {}
    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        kod = excel.metin(kayit.get("urun_kodu"))
        if not kod:
            sonuc.hatalar.append(SatirHatasi(satir_no, "-", "Ürün kodu boş olamaz"))
            continue

        degerler = {
            "_satir_no": satir_no,
            "urun_adi": excel.metin(kayit.get("urun_adi")),
            "urun_grubu": excel.metin(kayit.get("urun_grubu")),
            "aktif": excel.evet_hayir(kayit.get("aktif"), True),
        }
        for alan in sayisal_alanlar:
            degerler[alan] = excel.sayi_ya_da(kayit.get(alan))
        for alan in tam_sayi_alanlari:
            degerler[alan] = excel.tam_sayi_ya_da(kayit.get(alan))

        onceki = birlesik.get(kod)
        if onceki is None:
            birlesik[kod] = degerler
            continue
        sonuc.birlestirilen += 1
        for alan, deger in degerler.items():
            if alan == "_satir_no" or deger is None:
                continue
            if onceki.get(alan) is None:
                onceki[alan] = deger

    mevcutlar = {u.urun_kodu: u for u in db.scalars(select(IhracatUrunu)).all()}
    for kod, degerler in birlesik.items():
        urun = mevcutlar.get(kod)
        if urun is None:
            urun = IhracatUrunu(urun_kodu=kod)
            db.add(urun)
            mevcutlar[kod] = urun
            sonuc.eklenen += 1
        else:
            sonuc.guncellenen += 1

        urun.urun_adi = degerler.get("urun_adi") or ""
        urun.urun_grubu = degerler.get("urun_grubu")
        for alan in sayisal_alanlar + tam_sayi_alanlari:
            setattr(urun, alan, degerler.get(alan))
        urun.aktif = bool(degerler.get("aktif", True))

        if not urun.olculebilir_mi:
            sonuc.uyarilar.append(
                SatirHatasi(
                    degerler["_satir_no"], kod,
                    "Tır/konteyner yükleme adeti ve desi boş — doluluk hesaplanamaz",
                )
            )

    _aktarim_kaydet(db, dosya_adi, "IHRACAT_URUN", sonuc, kullanici)
    db.flush()
    return sonuc


def ihracat_musterilerini_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """İhracat müşteri master datasını yükler.

    Araç tipi burada belirlenir ve taşıma modunu da o belirler: konteyner yüklenen
    müşteri deniz, tır yüklenen kara yoludur. Sefer numarasının belge kodu (N/E) da
    müşteri bazındadır.
    """
    from app.domain.ihracat import arac_tipi_coz
    from app.models import IhracatMusterisi
    from app.services.veri_formatlari import (
        IHRACAT_MUSTERI_ALANLARI,
        IHRACAT_MUSTERI_ALIAS,
    )

    _kontrol_et(dosya, IHRACAT_MUSTERI_ALANLARI, IHRACAT_MUSTERI_ALIAS)
    kayitlar = excel.satirlari_oku(
        dosya, IHRACAT_MUSTERI_ALIAS, zorunlu_alanlar(IHRACAT_MUSTERI_ALANLARI)
    )
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))

    mevcutlar = {m.anahtar: m for m in db.scalars(select(IhracatMusterisi)).all()}
    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        ad = excel.metin(kayit.get("musteri_adi"))
        anahtar = yer_adi(ad)
        if not anahtar:
            sonuc.hatalar.append(SatirHatasi(satir_no, "-", "Müşteri adı boş olamaz"))
            continue

        musteri = mevcutlar.get(anahtar)
        if musteri is None:
            musteri = IhracatMusterisi(anahtar=anahtar, musteri_adi=ad)
            db.add(musteri)
            mevcutlar[anahtar] = musteri
            sonuc.eklenen += 1
        else:
            sonuc.guncellenen += 1

        musteri.musteri_adi = ad
        musteri.ulke = excel.metin(kayit.get("ulke")) or None
        musteri.ulke_kodu = (excel.metin(kayit.get("ulke_kodu")) or "").upper() or None
        musteri.sevk_adresi = excel.metin(kayit.get("sevk_adresi")) or None
        musteri.arac_tipi = arac_tipi_coz(excel.metin(kayit.get("arac_tipi"))).value
        kod = (excel.metin(kayit.get("sefer_kodu")) or "").upper()
        # NSC&Core sütunu doğrudan da verilebiliyor: 'NSC' -> N, 'Export' -> E.
        musteri.sefer_kodu = "N" if kod.startswith("N") else "E"
        musteri.yukleme_tipi = excel.metin(kayit.get("yukleme_tipi")) or None
        musteri.azami_agirlik = _tonaj(kayit.get("azami_agirlik"))
        musteri.aciklama = excel.metin(kayit.get("aciklama")) or None
        musteri.incoterms = (excel.metin(kayit.get("incoterms")) or "").upper() or None
        musteri.tedarikci = excel.metin(kayit.get("tedarikci")) or None
        musteri.satis_destek = excel.metin(kayit.get("satis_destek")) or None
        musteri.aktif = excel.evet_hayir(kayit.get("aktif"), True)

    _aktarim_kaydet(db, dosya_adi, "IHRACAT_MUSTERI", sonuc, kullanici)
    db.flush()
    return sonuc


def ihracat_siparislerini_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """İhracat siparişlerini yükler.

    İç piyasa dosyasından farkı: desi ve kg satır bazında dosyada gelir (ürün master
    datasında ihracat SKU'ları yok), müşteri adı `bayi_adi`, ülke `sehir` alanına
    yazılır. Satırlar IHRACAT havuzuna girer; diğer modüllerde görünmez.
    """
    from app.services.veri_formatlari import (
        IHRACAT_SIPARIS_ALANLARI,
        IHRACAT_SIPARIS_ALIAS,
    )

    _kontrol_et(dosya, IHRACAT_SIPARIS_ALANLARI, IHRACAT_SIPARIS_ALIAS)
    kayitlar = excel.satirlari_oku(
        dosya, IHRACAT_SIPARIS_ALIAS, zorunlu_alanlar(IHRACAT_SIPARIS_ALANLARI)
    )
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))
    aktarim = _aktarim_kaydet(db, dosya_adi, "SIPARIS/IHRACAT", sonuc, kullanici)
    db.flush()

    mevcutlar = _mevcut_siparis_satirlari(db, kayitlar)
    parti: dict[tuple[str, str, str], SiparisSatiri] = {}
    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        siparis_no = excel.metin(kayit.get("siparis_no"))
        urun_kodu = excel.metin(kayit.get("urun_kodu"))
        siparis_satir_no = urun_kodu
        anahtar = f"{siparis_no or '-'}/{urun_kodu or '-'}"
        try:
            if not siparis_no:
                raise ExcelHatasi("Sipariş No boş olamaz")
            if not urun_kodu:
                raise ExcelHatasi("Ürün kodu boş olamaz")
            teslimat_no = excel.metin(kayit.get("teslimat_no")) or siparis_no
            musteri_adi = excel.metin(kayit.get("bayi_adi"))
            if not musteri_adi:
                raise ExcelHatasi("Müşteri adı boş olamaz")
            miktar = excel.sayi(kayit.get("miktar"), "Adet")
            if miktar <= 0:
                raise ExcelHatasi("Adet sıfırdan büyük olmalı")
            desi = excel.sayi(kayit.get("desi"), "Desi")
            if desi <= 0:
                raise ExcelHatasi(
                    "Desi sıfırdan büyük olmalı; ihracatta araç kapasitesi desi ile ölçülür"
                )
            depo_kodu = excel.metin(kayit.get("depo_kodu"))
            if not depo_kodu:
                raise ExcelHatasi("Depo kodu boş olamaz")
        except ExcelHatasi as hata:
            sonuc.hatalar.append(SatirHatasi(satir_no, anahtar, str(hata)))
            continue

        satir_anahtari = (siparis_no, teslimat_no, siparis_satir_no)
        if satir_anahtari in parti:
            onceki = parti[satir_anahtari]
            onceki.miktar = Decimal(onceki.miktar) + miktar
            onceki.desi = Decimal(onceki.desi or 0) + desi
            onceki.agirlik = Decimal(onceki.agirlik or 0) + (
                excel.sayi_ya_da(kayit.get("agirlik")) or Decimal(0)
            )
            sonuc.birlestirilen += 1
            continue

        mevcut = mevcutlar.get(satir_anahtari)
        if mevcut is not None and mevcut.durum in {
            SiparisDurumu.PLANLANDI,
            SiparisDurumu.TAMAMLANDI,
        }:
            sonuc.atlanan += 1
            continue
        if mevcut is None:
            mevcut = SiparisSatiri(
                siparis_no=siparis_no,
                teslimat_no=teslimat_no,
                siparis_satir_no=siparis_satir_no,
            )
            db.add(mevcut)
            mevcutlar[satir_anahtari] = mevcut
            sonuc.eklenen += 1
        else:
            sonuc.guncellenen += 1

        mevcut.urun_kodu = urun_kodu
        mevcut.urun_adi = excel.metin(kayit.get("urun_adi"))
        mevcut.miktar = miktar
        mevcut.depo_kodu = depo_kodu.strip().upper()
        mevcut.bayi_adi = musteri_adi
        mevcut.sehir = excel.metin(kayit.get("sehir"))
        mevcut.ulke_kodu = (excel.metin(kayit.get("ulke_kodu")) or "").upper() or None
        mevcut.sevk_adresi = excel.metin(kayit.get("sevk_adresi"))
        mevcut.desi = desi
        mevcut.agirlik = excel.sayi_ya_da(kayit.get("agirlik"))
        mevcut.teslim_sekli = excel.metin(kayit.get("teslim_sekli"))
        mevcut.incoterms = (mevcut.teslim_sekli or "").upper() or None
        mevcut.siparis_tarihi = excel.tarih(kayit.get("siparis_tarihi"))
        mevcut.termin_tarihi = excel.tarih(kayit.get("termin_tarihi"))
        mevcut.durum = SiparisDurumu.BEKLEMEDE
        mevcut.hata_aciklamasi = None
        mevcut.bekleme_sebebi = None
        mevcut.modul = "IHRACAT"
        mevcut.ice_aktarim_id = aktarim.id
        parti[satir_anahtari] = mevcut

    db.flush()
    aktarim.basarili_satir = sonuc.basarili
    aktarim.hatali_satir = sonuc.hatali
    aktarim.hata_ozeti = _hata_ozeti(sonuc)
    db.flush()
    return sonuc


def urun_baglarini_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """Birlikte sevk edilecek ürün bağlarını topluca yükler.

    Bağ sayısı elle girilemeyecek kadar çok olabilir (2.585 ürünlük master datada
    kombi/baca eşleşmeleri yüzlerce satır tutar), bu yüzden Excel'den de yüklenir.
    Aynı çift tekrar gelirse tipi ve açıklaması güncellenir.
    """
    from app.services import urun_bagi_servisi

    _kontrol_et(dosya, URUN_BAGI_ALANLARI, URUN_BAGI_ALIAS)
    kayitlar = excel.satirlari_oku(
        dosya, URUN_BAGI_ALIAS, zorunlu_alanlar(URUN_BAGI_ALANLARI)
    )
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))

    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        ana = excel.metin(kayit.get("ana_urun_kodu"))
        bagli = excel.metin(kayit.get("bagli_urun_kodu"))
        tip = excel.metin(kayit.get("tip")).upper() or "AKSESUAR"
        try:
            oncesi = db.scalar(
                select(func.count(UrunBagi.id)).where(
                    UrunBagi.ana_urun_kodu.in_({ana, bagli}),
                    UrunBagi.bagli_urun_kodu.in_({ana, bagli}),
                )
            )
            urun_bagi_servisi.bag_kaydet(
                db, ana, bagli, tip,
                aciklama=excel.metin(kayit.get("aciklama")),
                kaynak="EXCEL",
            )
        except urun_bagi_servisi.BagHatasi as hata:
            sonuc.hatalar.append(SatirHatasi(satir_no, ana or "-", str(hata)))
            continue
        if oncesi:
            sonuc.guncellenen += 1
        else:
            sonuc.eklenen += 1

    db.commit()
    _aktarim_kaydet(db, dosya_adi, "URUN_BAGI", sonuc, kullanici)
    return sonuc


def tarifeleri_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """Şehir bazlı nakliye tarifelerini topluca yükler.

    Aynı anahtar (tip + il + araç + başlangıç + nakliyeci) tekrar gelirse fiyat
    güncellenir; yeni bir başlangıç tarihi yeni tarife satırı açar, eskisi
    geçmişteki planları maliyetlemek için durur.
    """
    from app.services import maliyet_servisi

    _kontrol_et(dosya, TARIFE_ALANLARI, TARIFE_ALIAS)
    kayitlar = excel.satirlari_oku(dosya, TARIFE_ALIAS, zorunlu_alanlar(TARIFE_ALANLARI))
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))

    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        il = excel.metin(kayit.get("il")) or ""
        try:
            baslangic = excel.tarih(kayit.get("gecerlilik_baslangic"))
            if baslangic is None:
                raise maliyet_servisi.MaliyetHatasi(
                    "Geçerlilik başlangıcı okunamadı (gg.aa.yyyy bekleniyor)."
                )
            oncesi = db.scalar(select(func.count(NakliyeTarifesi.id)))
            maliyet_servisi.tarife_kaydet(
                db,
                sevkiyat_tipi=excel.metin(kayit.get("sevkiyat_tipi")) or "",
                il=il,
                birim_fiyat=kayit.get("birim_fiyat"),
                gecerlilik_baslangic=baslangic,
                arac_tipi=excel.metin(kayit.get("arac_tipi")) or "",
                gecerlilik_bitis=excel.tarih(kayit.get("gecerlilik_bitis")),
                nakliyeci=excel.metin(kayit.get("nakliyeci")) or "",
                para_birimi=excel.metin(kayit.get("para_birimi")) or "TRY",
                aciklama=excel.metin(kayit.get("aciklama")) or "",
                ilce=excel.metin(kayit.get("ilce")) or "",
                cikis_noktasi=excel.metin(kayit.get("cikis_noktasi")) or "",
                desi_alt=kayit.get("desi_alt"),
                desi_ust=kayit.get("desi_ust"),
                motorin_fiyati=kayit.get("motorin_fiyati"),
            )
        except maliyet_servisi.MaliyetHatasi as hata:
            sonuc.hatalar.append(SatirHatasi(satir_no, il or "-", str(hata)))
            continue
        # Yeni satır mı yoksa mevcut fiyatın düzeltmesi mi: tablo büyüdüyse yeni.
        if db.scalar(select(func.count(NakliyeTarifesi.id))) > oncesi:
            sonuc.eklenen += 1
        else:
            sonuc.guncellenen += 1

    db.commit()
    _aktarim_kaydet(db, dosya_adi, "TARIFE", sonuc, kullanici)
    return sonuc


def butceyi_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """Aylık bütçe ve FC satırlarını yükler. Aynı ay/senaryo/sürüm tekrar gelirse
    tutar güncellenir — FC revizyonu böyle geçilir."""
    from app.services import maliyet_servisi

    _kontrol_et(dosya, BUTCE_ALANLARI, BUTCE_ALIAS)
    kayitlar = excel.satirlari_oku(dosya, BUTCE_ALIAS, zorunlu_alanlar(BUTCE_ALANLARI))
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))

    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        anahtar = f"{kayit.get('yil')}/{kayit.get('ay')}"
        try:
            maliyet_servisi.butce_kaydet(
                db,
                yil=kayit.get("yil"),
                ay=kayit.get("ay"),
                senaryo=excel.metin(kayit.get("senaryo")) or "",
                tutar=kayit.get("tutar"),
                surum=excel.metin(kayit.get("surum")) or "",
                marka=excel.metin(kayit.get("marka")) or "",
                desi=kayit.get("desi"),
                sevkiyat_tipi=excel.metin(kayit.get("sevkiyat_tipi")) or "",
                para_birimi=excel.metin(kayit.get("para_birimi")) or "TRY",
                aciklama=excel.metin(kayit.get("aciklama")) or "",
            )
            sonuc.eklenen += 1
        except maliyet_servisi.MaliyetHatasi as hata:
            sonuc.hatalar.append(SatirHatasi(satir_no, anahtar, str(hata)))

    db.commit()
    _aktarim_kaydet(db, dosya_adi, "BUTCE", sonuc, kullanici)
    return sonuc


def ek_ucretleri_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """Sözleşmenin sefer/desi dışı kalemlerini yükler: asgari gönderi bedeli,
    ek uğrama, ek km."""
    from app.services import maliyet_servisi

    _kontrol_et(dosya, EK_UCRET_ALANLARI, EK_UCRET_ALIAS)
    kayitlar = excel.satirlari_oku(
        dosya, EK_UCRET_ALIAS, zorunlu_alanlar(EK_UCRET_ALANLARI)
    )
    sonuc = IceAktarimSonucu(toplam=len(kayitlar))
    for kayit in kayitlar:
        satir_no = kayit["_satir_no"]
        tur = excel.metin(kayit.get("tur")) or ""
        try:
            baslangic = excel.tarih(kayit.get("gecerlilik_baslangic"))
            if baslangic is None:
                raise maliyet_servisi.MaliyetHatasi(
                    "Geçerlilik başlangıcı okunamadı (gg.aa.yyyy bekleniyor)."
                )
            maliyet_servisi.ek_ucret_kaydet(
                db, tur=tur, tutar=kayit.get("tutar"),
                gecerlilik_baslangic=baslangic,
                sevkiyat_tipi=excel.metin(kayit.get("sevkiyat_tipi")) or "",
                arac_tipi=excel.metin(kayit.get("arac_tipi")) or "",
                cikis_noktasi=excel.metin(kayit.get("cikis_noktasi")) or "",
                esik=kayit.get("esik"),
                gecerlilik_bitis=excel.tarih(kayit.get("gecerlilik_bitis")),
                nakliyeci=excel.metin(kayit.get("nakliyeci")) or "",
                motorin_fiyati=kayit.get("motorin_fiyati"),
                aciklama=excel.metin(kayit.get("aciklama")) or "",
            )
            sonuc.eklenen += 1
        except maliyet_servisi.MaliyetHatasi as hata:
            sonuc.hatalar.append(SatirHatasi(satir_no, tur or "-", str(hata)))

    db.commit()
    _aktarim_kaydet(db, dosya_adi, "EK_UCRET", sonuc, kullanici)
    return sonuc


def nakliye_tarifesini_aktar(
    db: Session, dosya: Path | Any, dosya_adi: str, kullanici: str = "sistem"
) -> IceAktarimSonucu:
    """Tarife ve ek ücretleri **tek dosyadan** yükler.

    Omsan listesi iki parçadan oluşuyor (fiyat tablosu + sözleşmenin ek
    kalemleri); ikisi ayrı sayfada gelir. 'Ek Ücretler' sayfası yoksa yalnızca
    tarifeler alınır.
    """
    from openpyxl import load_workbook

    sonuc = tarifeleri_aktar(db, dosya, dosya_adi, kullanici)
    try:
        if hasattr(dosya, "seek"):
            dosya.seek(0)
        kitap = load_workbook(dosya, read_only=True)
        sayfalar = set(kitap.sheetnames)
    except Exception:
        return sonuc
    if "Ek Ücretler" not in sayfalar:
        return sonuc
    if hasattr(dosya, "seek"):
        dosya.seek(0)
    ek = ek_ucretleri_aktar(db, dosya, dosya_adi, kullanici)
    sonuc.toplam += ek.toplam
    sonuc.eklenen += ek.eklenen
    sonuc.guncellenen += ek.guncellenen
    sonuc.hatalar.extend(ek.hatalar)
    return sonuc
