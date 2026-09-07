"""Geçmiş sevk planlarını sisteme alır: gerçekleşen maliyetin kaynağı.

Şirket sevkleri yıl boyunca tek bir "Sevk Planları" dosyasında tutuyor: her satır
bir sipariş kalemi, **Belge No** ise aracın sefer numarası. Maliyet modülünün
"gerçekleşen" sütunu bu dosyadan doluyor — plan üretmeden geçmişi maliyetlemenin
başka yolu yok.

**Belge kodu sevkiyat tipini söyler.** Numaranın harf bölümü (2601**S**2001) hangi
tip olduğunu verir ve maliyetin hesaplanıp hesaplanmayacağını da o belirler:

| Kod | Anlamı | Maliyet |
|---|---|---|
| `D` | Ring | **Alınmaz** — ring kendi deposundan çıkar, nakliyeciye sefer bedeli ödenmez |
| `S` | FTL (tam araç) | Sefer fiyatı |
| `R` | Rutin / parsiyel | Kademeli desi fiyatı |
| `K` | Kargo | Desi fiyatı |
| `B` | Arçelik / bayi dağıtım deposu | Tam araç sayılır, FTL fiyatı |
| `A` | Alıcı vasıtası | **Yok** — taşımayı müşteri üstleniyor |
| `T` | Sistemsel | **Yok** — fiziksel araç hareketi değil |
| `ST` | Stok aktarımı | **Hesaplanmaz** — depolar arası shuttle, ayrı fiyatlanıyor |

Maliyeti olmayan tipler `sevkiyat_tipi` boş bırakılarak alınır: maliyet motoru
onları sıfır sayar ve "tarifesi eksik" diye **işaretlemez**. İkisini karıştırmak
raporu yanıltırdı — bedeli olmayan sevk ile fiyatı bilinmeyen sevk ayrı şeyler.
"""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.iller import yer_adi
from app.domain.marka import marka as depo_markasi
from app.models import PlanDurumu, SevkiyatPlani, SiparisDurumu, SiparisSatiri

# Sevk dosyalarında kolon yerleşimi sabit (bkz. scripts/ic_piyasa_analiz.py).
ARAC_TIPI, SEHIR, SIPARIS_NO, BELGE_NO, DEPO = 3, 4, 5, 6, 7
STOK_KODU, STOK_ADI, ADET, BAYI, ALICI, ADRES, NOT_ALANI = 8, 9, 10, 11, 12, 13, 14
TARIH, MUSTERI_KODU, TESLIMAT, PLAN_TARIHI, DAGITIM, DESI = 15, 16, 17, 18, 19, 20

BELGE_DESENI = re.compile(r"^(\d{2})(\d{2})([A-ZÇŞİĞÜÖ]+)(\d+)")

RING_KODU = "D"
BELGE_TIPLERI: dict[str, str | None] = {
    "S": "FTL",
    "R": "RUTIN",
    "K": "KARGO",
    "B": "FTL",
    "A": None,
    "T": None,
    "ST": None,
}
"""Belge kodu -> sevkiyat tipi. `None` = maliyeti olmayan hareket."""

MALIYETSIZ_ACIKLAMALARI = {
    "A": "Alıcı vasıtası — taşımayı müşteri üstlendi, navlun ödenmiyor.",
    "T": "Sistemsel hareket — fiziksel araç yok.",
    "ST": "Stok aktarımı — depolar arası shuttle, ayrı fiyatlanıyor.",
}

ARAC_ESLEMESI = (
    ("KAMYON", "KAMYON"),
    ("TIR", "TIR"),
    ("TİR", "TIR"),
)
"""Serbest metin araç alanından araç tipi. Alan sahada elle doldurulduğu için
'RUTİN KAMYON', 'FTL KAMYON', 'CEVA RUTİN' gibi çok sayıda varyant içeriyor;
içinde KAMYON geçen kamyon, TIR geçen tırdır."""


class SevkAktarimHatasi(Exception):
    """Ekranda gösterilecek, kullanıcının düzeltebileceği hata."""


@dataclass
class AktarimOzeti:
    okunan_satir: int = 0
    plan: int = 0
    ring_atlanan: int = 0
    maliyetsiz: int = 0
    guncellenen: int = 0
    satir: int = 0
    bozuk_desi: int = 0
    tanimsiz_kod: dict[str, int] = field(default_factory=dict)
    nakliyeciler: dict[str, int] = field(default_factory=dict)

    def ozet(self) -> str:
        parcalar = [
            f"{self.okunan_satir} satır okundu",
            f"{self.plan} sefer alındı ({self.satir} sipariş satırı)",
            f"{self.ring_atlanan} ring satırı atlandı",
        ]
        if self.guncellenen:
            parcalar.append(f"{self.guncellenen} sefer güncellendi")
        if self.maliyetsiz:
            parcalar.append(f"{self.maliyetsiz} sefer maliyetsiz (alıcı vasıtası, "
                            "sistemsel, stok aktarımı)")
        if self.bozuk_desi:
            parcalar.append(f"{self.bozuk_desi} satırda desi okunamadı")
        if self.tanimsiz_kod:
            bilinmeyen = ", ".join(
                f"{kod} ({adet})" for kod, adet in sorted(self.tanimsiz_kod.items())
            )
            parcalar.append(f"tanınmayan belge kodu: {bilinmeyen}")
        return " · ".join(parcalar)


def _metin(deger: Any) -> str:
    return str(deger).strip() if deger is not None else ""


def _ondalik(deger: Any) -> Decimal:
    """Sayı olmayan hücre sıfır sayılır.

    Kaynak dosyada formül artığı `#REF!` hücreleri var; bunlar için aktarımı
    durdurmak yerine sıfır yazıp sayısını raporluyoruz.
    """
    if deger is None or isinstance(deger, str) and deger.startswith("#"):
        return Decimal(0)
    try:
        return Decimal(str(deger))
    except (InvalidOperation, ValueError):
        return Decimal(0)


def _tarih(deger: Any) -> date | None:
    if isinstance(deger, datetime):
        return deger.date()
    if isinstance(deger, date):
        return deger
    return None


def _teslimat_anahtari(teslimat_no: str, siparis_no: str) -> str:
    """Teslimatın bölünemez anahtarı.

    Bayi ortak deposu (-1) satırlarında bu sütun teslimat numarası yerine
    "BAYİ DEPO" gibi bir etiket taşıyor; sipariş bölünemez birim olduğu için
    anahtar sipariş numarasından kurulur (bkz. ice_aktarim, aynı kural).
    """
    if not teslimat_no:
        return siparis_no
    if not any(karakter.isdigit() for karakter in teslimat_no):
        return f"{siparis_no}-{yer_adi(teslimat_no) or 'SIPARIS'}"
    return teslimat_no


def _arac_tipi(deger: str) -> str | None:
    ad = yer_adi(deger)
    for desen, tip in ARAC_ESLEMESI:
        if desen in ad:
            return tip
    return None


def belge_coz(belge: str) -> tuple[str, str, str] | None:
    """Belge numarasından (dönem, kod, sefer no). Tanınmayan numara için None."""
    eslesme = BELGE_DESENI.match(belge)
    if not eslesme:
        return None
    donem = eslesme.group(1) + eslesme.group(2)
    return donem, eslesme.group(3), belge


@dataclass
class _Sefer:
    """Dosyadan okunan ham sefer; veritabanına yazılmadan önce toplanır."""

    belge_no: str
    donem: str
    kod: str
    satirlar: list[dict] = field(default_factory=list)
    arac_tipleri: list[str] = field(default_factory=list)
    nakliyeciler: list[str] = field(default_factory=list)
    plan_tarihi: date | None = None

    @property
    def tipi(self) -> str | None:
        return BELGE_TIPLERI.get(self.kod)

    def _sirali_iller(self) -> list[str]:
        """Uğranan iller, **yakından uzağa**.

        Maliyet motoru FTL sefer fiyatını listenin sonuncusundan (en uzak il)
        okuyor; sıralama tesise uzaklığa göre kurulur.
        """
        from app.domain.iller import ESKISEHIR_MESAFELERI

        iller = {s["il"] for s in self.satirlar if s["il"]}
        return sorted(iller, key=lambda il: (ESKISEHIR_MESAFELERI.get(il, 9999), il))


def _satirlari_topla(sayfa, ozet: AktarimOzeti) -> dict[str, _Sefer]:
    seferler: dict[str, _Sefer] = {}
    for satir in sayfa.iter_rows(min_row=2, values_only=True):
        if len(satir) <= DESI:
            continue
        belge = _metin(satir[BELGE_NO])
        if not belge:
            continue
        ozet.okunan_satir += 1

        cozum = belge_coz(belge)
        if cozum is None:
            ozet.tanimsiz_kod["?"] = ozet.tanimsiz_kod.get("?", 0) + 1
            continue
        donem, kod, sefer_no = cozum
        if kod == RING_KODU:
            ozet.ring_atlanan += 1
            continue
        if kod not in BELGE_TIPLERI:
            ozet.tanimsiz_kod[kod] = ozet.tanimsiz_kod.get(kod, 0) + 1
            continue

        sefer = seferler.setdefault(sefer_no, _Sefer(sefer_no, donem, kod))
        desi_hucresi = satir[DESI]
        if isinstance(desi_hucresi, str) and desi_hucresi.startswith("#"):
            ozet.bozuk_desi += 1

        siparis_no = _metin(satir[SIPARIS_NO])
        sefer.satirlar.append({
            "siparis_no": siparis_no,
            "teslimat_no": _teslimat_anahtari(_metin(satir[TESLIMAT]), siparis_no),
            "urun_kodu": _metin(satir[STOK_KODU]),
            "urun_adi": _metin(satir[STOK_ADI]),
            "miktar": _ondalik(satir[ADET]),
            "desi": _ondalik(desi_hucresi),
            "depo_kodu": _metin(satir[DEPO]).replace(" ", ""),
            "il": yer_adi(satir[SEHIR]),
            "sehir": _metin(satir[SEHIR]),
            "ilce": _metin(satir[NOT_ALANI]).lstrip(" -").strip(),
            "bayi_adi": _metin(satir[BAYI]),
            "alici_firma": _metin(satir[ALICI]),
            "sevk_adresi": _metin(satir[ADRES]),
            "tarih": _tarih(satir[TARIH]),
        })
        arac = _arac_tipi(_metin(satir[ARAC_TIPI]))
        if arac:
            sefer.arac_tipleri.append(arac)
        nakliyeci = _metin(satir[DAGITIM])
        if nakliyeci:
            sefer.nakliyeciler.append(nakliyeci)
        if sefer.plan_tarihi is None:
            sefer.plan_tarihi = _tarih(satir[PLAN_TARIHI])
    return seferler


def _cogunluk(degerler: list[str]) -> str | None:
    if not degerler:
        return None
    sayac: dict[str, int] = defaultdict(int)
    for deger in degerler:
        sayac[deger] += 1
    return max(sayac.items(), key=lambda ikili: (ikili[1], ikili[0]))[0]


def _plani_yaz(db: Session, sefer: _Sefer, ozet: AktarimOzeti) -> SevkiyatPlani:
    """Bir seferi plan + sipariş satırı olarak yazar; varsa günceller."""
    mevcut = db.scalar(
        select(SevkiyatPlani).where(SevkiyatPlani.sefer_no == sefer.belge_no)
    )
    if mevcut is not None:
        # Aynı dosya tekrar yüklenebilir: sefer baştan kurulur, iki kez sayılmaz.
        for satir in list(mevcut.satirlar):
            db.delete(satir)
        plan = mevcut
        ozet.guncellenen += 1
    else:
        plan = SevkiyatPlani(sefer_no=sefer.belge_no)
        db.add(plan)

    iller = sefer._sirali_iller()
    duraklar = {
        (s["bayi_adi"] or s["alici_firma"], s["il"], s["ilce"]) for s in sefer.satirlar
    }
    musteriler = {s["bayi_adi"] or s["alici_firma"] for s in sefer.satirlar}
    toplam_desi = sum((s["desi"] for s in sefer.satirlar), Decimal(0))
    toplam_adet = sum((s["miktar"] for s in sefer.satirlar), Decimal(0))
    depolar = [s["depo_kodu"] for s in sefer.satirlar if s["depo_kodu"]]

    plan.donem = sefer.donem
    plan.modul = "ROTA"
    plan.plan_tipi = "IC_FTL"
    plan.depo_kodu = _cogunluk(depolar) or "64"
    plan.planlama_anahtari = sefer.kod
    plan.urun_kodlari = ",".join(
        sorted({s["urun_kodu"] for s in sefer.satirlar if s["urun_kodu"]})
    )[:500]
    plan.olcu = "ANAHTAR"
    plan.toplam_birim = Decimal(0)
    plan.toplam_adet = toplam_adet
    plan.toplam_desi = toplam_desi
    plan.doluluk_yuzdesi = Decimal(0)
    plan.teslimat_sayisi = len({s["teslimat_no"] for s in sefer.satirlar})
    plan.plan_tarihi = sefer.plan_tarihi or date(
        2000 + int(sefer.donem[:2]), int(sefer.donem[2:]), 1
    )
    plan.sevkiyat_tipi = sefer.tipi
    plan.arac_tipi = _cogunluk(sefer.arac_tipleri)
    plan.nakliyeci = _cogunluk(sefer.nakliyeciler)
    plan.durak_sayisi = len(duraklar)
    plan.musteri_sayisi = len(musteriler)
    plan.iller_metni = ", ".join(iller)[:400]
    plan.ilceler_metni = "+".join(
        sorted({s["ilce"] for s in sefer.satirlar if s["ilce"]})
    )[:600]
    plan.son_ugrak = iller[-1] if iller else None
    plan.durum = PlanDurumu.TAMAMLANDI
    plan.olusturan = "sevk aktarımı"
    plan.marka_paylari_metni = _marka_paylari(sefer)
    if sefer.tipi is None:
        plan.maliyet_notu = MALIYETSIZ_ACIKLAMALARI.get(sefer.kod)
        ozet.maliyetsiz += 1
    db.flush()

    for sira, satir in enumerate(sefer.satirlar, start=1):
        db.add(SiparisSatiri(
            siparis_no=satir["siparis_no"] or sefer.belge_no,
            # Satır numarası sefer adıyla benzersizleşir: aynı sipariş/teslimat
            # ikilisi farklı seferlerde tekrar geçebiliyor.
            siparis_satir_no=f"{sefer.belge_no}-{sira}"[:20],
            teslimat_no=satir["teslimat_no"] or f"{sefer.belge_no}-{sira}",
            urun_kodu=satir["urun_kodu"],
            urun_adi=satir["urun_adi"][:250] or None,
            miktar=satir["miktar"],
            depo_kodu=satir["depo_kodu"] or plan.depo_kodu,
            sehir=satir["sehir"][:80] or None,
            ilce=satir["ilce"][:80] or None,
            bayi_adi=satir["bayi_adi"][:250] or None,
            alici_firma=satir["alici_firma"][:250] or None,
            sevk_adresi=satir["sevk_adresi"][:250] or None,
            desi=satir["desi"],
            siparis_tarihi=satir["tarih"],
            modul="ROTA",
            plan_id=plan.id,
            durum=SiparisDurumu.TAMAMLANDI,
        ))
        ozet.satir += 1
    ozet.plan += 1
    return plan


def _marka_paylari(sefer: _Sefer) -> str | None:
    """Seferin markalar arası dağılımı; navlun faturası marka bazında kesiliyor."""
    from app.domain.marka import paylari_hesapla, paylari_metne_cevir

    katkilar: dict[str, Decimal] = defaultdict(Decimal)
    for satir in sefer.satirlar:
        if satir["depo_kodu"]:
            katkilar[satir["depo_kodu"]] += satir["desi"] or satir["miktar"]
    if not katkilar:
        return None
    return paylari_metne_cevir(paylari_hesapla(dict(katkilar))) or None


def aktar(
    db: Session,
    dosya: Path | Any,
    sayfa_adi: str | None = None,
    yil: int | None = None,
) -> AktarimOzeti:
    """Sevk planları dosyasını okur ve seferleri sisteme yazar.

    Ring seferleri **alınmaz**: ring kendi deposundan çıkan bir iç sevkiyattır,
    nakliyeciye sefer bedeli ödenmez.
    """
    kitap = load_workbook(dosya, read_only=True, data_only=True)
    try:
        sayfa = kitap[sayfa_adi] if sayfa_adi else kitap.worksheets[0]
        ozet = AktarimOzeti()
        seferler = _satirlari_topla(sayfa, ozet)
    finally:
        kitap.close()

    if not seferler:
        raise SevkAktarimHatasi(
            "Dosyada ring dışı sefer bulunamadı. Belge No sütununun dolu olduğundan "
            "ve numaraların 2601S2001 biçiminde olduğundan emin olun."
        )

    for sefer in seferler.values():
        if yil and sefer.donem and int(sefer.donem[:2]) != yil % 100:
            continue
        _plani_yaz(db, sefer, ozet)
        for nakliyeci in set(sefer.nakliyeciler):
            ozet.nakliyeciler[nakliyeci] = ozet.nakliyeciler.get(nakliyeci, 0) + 1
    return ozet
